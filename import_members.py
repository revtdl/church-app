#!/usr/bin/env python3
"""
Direct Member Import Script for Church Nexus
Reads Excel file and imports members to Firebase
"""

import openpyxl
import requests
import json
import time
from datetime import datetime

# Firebase config
FIREBASE_PROJECT = "churchapp-74dec"
FIREBASE_API_KEY = "AIzaSyDXN2CFtyuZQbGKJbciB8Nm9WYTGVxXQ6g"
FIRESTORE_URL = f"https://firestore.googleapis.com/v1/projects/{FIREBASE_PROJECT}/databases/(default)/documents/members"

# Excel file path
EXCEL_FILE = r"C:\Users\amaze\Downloads\WMC_Membership_List_Final_2026.xlsx"

def get_existing_members():
    """Fetch existing members from Firestore"""
    try:
        url = f"https://firestore.googleapis.com/v1/projects/{FIREBASE_PROJECT}/databases/(default)/documents/members"
        response = requests.get(url, params={"key": FIREBASE_API_KEY})

        if response.status_code == 200:
            data = response.json()
            existing = set()
            if 'documents' in data:
                for doc in data['documents']:
                    fields = doc.get('fields', {})
                    name = fields.get('name', {}).get('stringValue', '').lower()
                    if name:
                        existing.add(name)
            return existing
        return set()
    except Exception as e:
        print(f"⚠️  Could not fetch existing members: {e}")
        return set()

def read_excel(filepath):
    """Read Excel file and extract members"""
    try:
        print("📖 Reading Excel file...")
        wb = openpyxl.load_workbook(filepath)

        # Find "All Members" sheet
        sheet_name = None
        for name in wb.sheetnames:
            if 'All Members' in name:
                sheet_name = name
                break

        if not sheet_name:
            raise Exception('Sheet "All Members" not found')

        ws = wb[sheet_name]
        print(f"✅ Found sheet: {sheet_name}")

        # Find header row
        header_row = None
        for row_idx, row in enumerate(ws.iter_rows(max_row=20, values_only=True), 1):
            if row[0] == 'Category':
                header_row = row_idx
                break

        if not header_row:
            raise Exception('Could not find header row')

        print(f"📋 Header row at line {header_row}\n")

        # Extract members
        members = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=header_row+1, values_only=True), header_row+1):
            try:
                # Column 3 is Full Name (index 3)
                name = str(row[3]).strip() if row[3] else None

                if not name or len(name) < 2 or name == 'None':
                    continue

                # Column mapping (0-indexed):
                # 0: Category, 1: S.No, 2: Family#, 3: Full Name, 4: Gender
                # 5: Head/Dependent, 6: Father, 7: Spouse, 8: DOB, 9: Month
                # 10: POB, 11: Wedding, 12: Wedding Month, 13: Baptised, 14: Baptism Date

                # Safe column access
                def safe_col(idx):
                    try:
                        val = row[idx] if idx < len(row) else None
                        return str(val).strip() if val and str(val) != 'None' else ''
                    except:
                        return ''

                member = {
                    'name': name,
                    'email': '',  # Not in Excel
                    'phone': '',  # Not in Excel
                    'phone2': '',
                    'familyHead': 'Head' if safe_col(5) == 'Head of Family' else 'Dependent',
                    'familyNumber': safe_col(2),  # Family #
                    'area': '',  # Not in Excel
                    'address': '',  # Not in Excel
                    'dateOfBirth': safe_col(8),  # Date of Birth (col 9 is 0-indexed as 8)
                    'maritalStatus': '',  # Not in Excel
                    'spouseName': safe_col(7),  # Spouse Name (col 8)
                    'membershipStatus': safe_col(0),  # Category (col 1)
                    'remarks': f"Gender: {safe_col(4)}, Place of Birth: {safe_col(10)}",
                    'profileImage': None,
                }

                members.append(member)
            except Exception as e:
                print(f"⚠️  Skipping row {row_idx}: {e}")
                continue

        return members
    except Exception as e:
        print(f"❌ Error reading Excel: {e}")
        raise

def import_members(members):
    """Import members to Firebase"""
    print(f"🚀 Importing {len(members)} members to Firebase...\n")

    # Get existing members
    existing = get_existing_members()
    print(f"🔍 Found {len(existing)} existing members\n")

    imported = 0
    duplicates = 0
    failed = 0

    for idx, member in enumerate(members, 1):
        try:
            # Check for duplicates
            if member['name'].lower() in existing:
                duplicates += 1
                continue

            # Prepare document
            doc_data = {
                "fields": {
                    "name": {"stringValue": member['name']},
                    "email": {"stringValue": member['email']},
                    "phone": {"stringValue": member['phone']},
                    "phone2": {"stringValue": member['phone2']},
                    "familyHead": {"stringValue": member['familyHead']},
                    "familyNumber": {"stringValue": member['familyNumber']},
                    "area": {"stringValue": member['area']},
                    "address": {"stringValue": member['address']},
                    "dateOfBirth": {"stringValue": member['dateOfBirth']},
                    "maritalStatus": {"stringValue": member['maritalStatus']},
                    "spouseName": {"stringValue": member['spouseName']},
                    "membershipStatus": {"stringValue": member['membershipStatus']},
                    "createdAt": {"timestampValue": datetime.utcnow().isoformat() + "Z"},
                }
            }

            # Import to Firestore
            response = requests.post(
                FIRESTORE_URL,
                params={"key": FIREBASE_API_KEY},
                json=doc_data
            )

            if response.status_code in [200, 201]:
                imported += 1
                percent = int((idx / len(members)) * 100)
                print(f"  [{percent}%] ✅ {member['name']}")
            else:
                failed += 1
                print(f"  ❌ Failed: {member['name']} - {response.status_code}")

            # Rate limiting
            if idx % 10 == 0:
                time.sleep(0.5)

        except Exception as e:
            failed += 1
            print(f"  ❌ Error importing {member['name']}: {e}")

    print(f"\n✅ SUCCESS!\n")
    print(f"📊 Summary:")
    print(f"   • Imported: {imported}")
    print(f"   • Duplicates skipped: {duplicates}")
    print(f"   • Failed: {failed}")
    print(f"   • Total: {len(existing) + imported} members\n")
    print(f"🌐 View members: http://localhost:8000/admin-portal.html\n")

if __name__ == '__main__':
    try:
        print("=" * 50)
        print("🏦 Church Nexus - Bulk Member Import")
        print("=" * 50 + "\n")

        members = read_excel(EXCEL_FILE)
        print(f"✅ Loaded {len(members)} members from Excel\n")

        import_members(members)

    except Exception as e:
        print(f"\n❌ Error: {e}")
